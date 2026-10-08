import io
import json
from pathlib import Path
from types import SimpleNamespace
import zipfile
import pytest
from assurance.ai import extract_records, explain
from assurance.documents import read_document, search_evidence
from assurance.engine import reconcile
from assurance.storage import Store, export_bundle, csv_bytes
from assurance.synthetic import generate

def test_persist_case_and_export(tmp_path):
    rules = json.loads(Path('examples/rules.json').read_text())
    rows,_ = generate(rules,20)
    store = Store(tmp_path/'db.sqlite3')
    run_id = store.save_run('Test','synthetic',rows,rules,reconcile(rows,rules))
    store.save_case(run_id,2,'Investigating','Analyst','Check notice')
    with pytest.raises(ValueError,match='rationale'):
        store.save_case(run_id,2,'Resolved','Analyst','')
    store.save_case(run_id,2,'Resolved','Analyst','Confirmed source evidence')
    reopened = Store(tmp_path/'db.sqlite3')
    assert reopened.cases(run_id)[2]['status'] == 'Resolved'
    assert len(reopened.audit(run_id)) == 3
    archive = zipfile.ZipFile(io.BytesIO(export_bundle(reopened,run_id)))
    assert {'rules.json','audit.json','reconciliation.csv','evidence.json','run.json','report.md'} == set(archive.namelist())
    assert json.loads(archive.read('run.json'))['rows'] == rows

def test_export_formula_protection():
    assert b"'=HYPERLINK" in csv_bytes([{'policy':'=HYPERLINK("malicious")'}])

def test_evidence_and_locator():
    doc = read_document('notice.txt',b'Policy P123. Premium AUD 1017.50.')
    assert search_evidence([doc],'P123')[0]['locator'] == 'text'
    assert len(doc['sha256']) == 64

class Client:
    def __init__(self,text):
        self.text = text
        self.responses = self
        self.kwargs = None
    def create(self,**kwargs):
        self.kwargs = kwargs
        return SimpleNamespace(output_text=self.text,id='mock-response',status='completed')

def test_ai_extraction_citations_and_no_storage(monkeypatch):
    monkeypatch.setenv('OPENAI_MODEL','test-model')
    doc = read_document('notice.txt',b'Policy P123. Premium AUD 1017.50.')
    record = {'policy_id':'P123','renewal_date':'','segment':'','risk_factor':'','discount_rate':'',
              'recorded_premium':'1017.50','applied_rule_version':'','source_document':'notice.txt',
              'source_locator':'text','evidence_quote':'Policy P123.'}
    client = Client(json.dumps({'records':[record],'warnings':['Missing rating fields']}))
    parsed,_ = extract_records([doc],client)
    assert parsed['records'][0]['risk_factor'] == ''
    assert client.kwargs['store'] is False
    record['evidence_quote'] = 'Invented evidence'
    client.text = json.dumps({'records':[record],'warnings':[]})
    with pytest.raises(ValueError,match='unverified citation'):
        extract_records([doc],client)

def test_ai_invalid_json(monkeypatch):
    monkeypatch.setenv('OPENAI_MODEL','test-model')
    with pytest.raises(ValueError,match='valid JSON'):
        extract_records([],Client('not JSON'))

def test_docx_and_pdf_text():
    from docx import Document
    from pypdf import PdfWriter
    doc = Document()
    doc.add_paragraph('Policy REAL-1')
    table = doc.add_table(rows=1,cols=2)
    table.cell(0,0).text = 'Premium'
    table.cell(0,1).text = '1017.50'
    stream = io.BytesIO()
    doc.save(stream)
    evidence = read_document('notice.docx',stream.getvalue())
    assert '1017.50' in evidence['chunks'][0]['text']
    writer = PdfWriter()
    writer.add_blank_page(width=400,height=600)
    stream = io.BytesIO()
    writer.write(stream)
    assert read_document('scan.pdf',stream.getvalue())['empty_pages'] == ['page 1']

def test_real_excel_date_cells_and_custom_column_mapping():
    from datetime import datetime
    from openpyxl import Workbook
    from assurance.ingestion import read_portfolio, map_records
    wb = Workbook()
    sheet = wb.active
    sheet.append(['ID','Date','Area','Risk','Discount','Premium'])
    sheet.append(['000123',datetime(2026,6,15),'Metro','1','0.1','1017.50'])
    stream = io.BytesIO()
    wb.save(stream)
    frame = read_portfolio(stream.getvalue(),'portfolio.xlsx')
    mapping = dict(policy_id='ID',renewal_date='Date',segment='Area',risk_factor='Risk',
                   discount_rate='Discount',recorded_premium='Premium',applied_rule_version='absent')
    rows = map_records(frame,mapping)
    rules = json.loads(Path('examples/rules.json').read_text())
    assert rows[0]['policy_id'] == '000123'
    assert rows[0]['renewal_date'] == '2026-06-15'
    assert reconcile(rows,rules)[0]['status'] == 'matched'
