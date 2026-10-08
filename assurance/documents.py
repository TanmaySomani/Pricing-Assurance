import hashlib
import io
from pathlib import Path


MAX_BYTES = 20 * 1024 * 1024


def read_document(name, data):
    if len(data) > MAX_BYTES:
        raise ValueError('Document exceeds the 20 MB local limit')
    suffix = Path(name).suffix.lower()
    chunks = []
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('Encrypted PDF: supply an unencrypted copy')
        if len(reader.pages) > 150:
            raise ValueError('PDF exceeds 150 pages; split into smaller documents')
        for page, content in enumerate(reader.pages, 1):
            text = content.extract_text() or ''
            chunks.append({'locator': f'page {page}', 'text': text})
    elif suffix == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        chunks.append({'locator':'document body and tables', 'text':text})
    elif suffix in ['.txt', '.md']:
        chunks.append({'locator':'text', 'text':data.decode('utf-8-sig')})
    else:
        raise ValueError('Supported evidence types: PDF, DOCX, TXT, MD')
    if sum(len(c['text']) for c in chunks) > 600000:
        raise ValueError('Extracted text exceeds 600,000 characters; split the document')
    return {'name':Path(name).name, 'sha256':hashlib.sha256(data).hexdigest(), 'chunks':chunks,
            'empty_pages':[c['locator'] for c in chunks if not c['text'].strip()]}


def evidence_text(documents, limit=40000):
    full = '\n\n'.join(f'[{d["name"]} | {c["locator"]}]\n{c["text"]}' for d in documents for c in d['chunks'])
    if len(full) > limit:
        raise ValueError(f'Selected evidence exceeds {limit:,} characters. Select fewer or smaller documents.')
    return full


def search_evidence(documents, query, limit=5):
    words = set(query.lower().split())
    candidates = []
    for doc in documents:
        for chunk in doc['chunks']:
            # Windows retain exact text and page locators; no invented citations.
            for offset in range(0, len(chunk['text']), 1800):
                text = chunk['text'][offset:offset+2200]
                score = sum(w in text.lower() for w in words)
                if score:
                    candidates.append({'name':doc['name'], 'sha256':doc['sha256'],
                                       'locator':chunk['locator'], 'text':text, 'score':score})
    return sorted(candidates, key=lambda c:c['score'], reverse=True)[:limit]
